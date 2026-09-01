"""
Testes de agrupamento, associacao e do pipeline completo.

O teste de ponta a ponta roda sobre a prancha de exemplo real e trava os
numeros conferidos manualmente contra a tabela de resumo do proprio DWG.
"""
import math

import pytest
from shapely.geometry import LineString

from extrator.agrupamento import agrupar, estatisticas_grupos
from extrator.associacao import AssociadorGeometrico, ponto_representativo
from extrator.modelos import GeometriaDXF, TextoDXF


def _texto(x, y, conteudo="10N1-Ø10c/15 C=400", rot=0.0, altura=0.2):
    return TextoDXF(handle="A", conteudo=conteudo, conteudo_bruto=conteudo,
                    x=x, y=y, altura=altura, rotacao=rot, layer="ARR",
                    prancha="P1")


def _barra(coords, handle="B", angulo=0.0, distribuicao=False):
    linha = LineString(coords)
    return GeometriaDXF(handle=handle, geometria=linha,
                        comprimento_desenho=linha.length, angulo=angulo,
                        layer="ARR", linetype="CONTINUOUS", prancha="P1",
                        e_distribuicao=distribuicao)


# =============================================================================
# Agrupamento de textos empilhados
# =============================================================================

def test_textos_no_mesmo_ponto_formam_um_grupo(cfg):
    cfg.calibrar(0.2)                       # tolerancia = 3 x 0.2 = 0.6
    textos = [_texto(10.0, 10.0), _texto(10.0, 10.3), _texto(10.0, 10.6)]
    grupos = agrupar(textos, cfg)
    assert len(grupos) == 1 and len(grupos[0].indices) == 3


def test_textos_distantes_ficam_em_grupos_separados(cfg):
    cfg.calibrar(0.2)
    textos = [_texto(10.0, 10.0), _texto(50.0, 10.0)]
    assert len(agrupar(textos, cfg)) == 2


def test_empilhados_continuam_sendo_posicoes_distintas(cfg):
    """O agrupamento nao funde textos: so marca que disputam barras vizinhas."""
    cfg.calibrar(0.2)
    textos = [_texto(10.0, 10.0, "10N1-Ø10c/15 C=400"),
              _texto(10.0, 10.3, "20N2-Ø12.5c/20 C=300")]
    grupos = agrupar(textos, cfg)
    assert len(grupos) == 1
    assert sorted(grupos[0].indices) == [0, 1]      # dois textos preservados


def test_agrupamento_de_muitos_textos_e_rapido(cfg):
    """Grid hashing: 5000 textos nao podem virar 12,5 milhoes de comparacoes."""
    cfg.calibrar(0.2)
    textos = [_texto(i * 5.0, j * 5.0) for i in range(70) for j in range(70)]
    grupos = agrupar(textos, cfg)
    assert len(grupos) == len(textos)            # todos isolados
    assert estatisticas_grupos(grupos) == {1: len(textos)}


# =============================================================================
# Associacao texto <-> barra
# =============================================================================

def test_texto_associa_a_barra_mais_proxima(cfg):
    cfg.calibrar(0.2)                       # raio = 12 x 0.2 = 2.4
    barras = [_barra([(0, 0), (0, 10)], "PERTO"),
              _barra([(20, 0), (20, 10)], "LONGE")]
    a = AssociadorGeometrico(barras, cfg)
    grupo = agrupar([_texto(0.3, 5.0, rot=90.0)], cfg)[0]
    res = a.associar_grupo(grupo, [_texto(0.3, 5.0, rot=90.0)])
    assert a.barras[res[0].indice_barra].handle == "PERTO"


def test_alinhamento_angular_desempata_barras_equidistantes(cfg):
    """O texto de armadura e escrito paralelo a sua barra."""
    cfg.calibrar(0.2)
    barras = [_barra([(-1, 5), (1, 5)], "HORIZONTAL", angulo=0.0),
              _barra([(0, 4), (0, 6)], "VERTICAL", angulo=90.0)]
    a = AssociadorGeometrico(barras, cfg)
    textos = [_texto(0.0, 5.0, rot=90.0)]        # texto vertical
    res = a.associar_grupo(agrupar(textos, cfg)[0], textos)
    assert a.barras[res[0].indice_barra].handle == "VERTICAL"


def test_dois_textos_empilhados_ocupam_duas_barras_diferentes(cfg):
    """Regra das barras empilhadas: atribuicao exclusiva dentro do grupo."""
    cfg.calibrar(0.2)
    barras = [_barra([(0, 0), (0, 10)], "B1", angulo=90.0),
              _barra([(0.4, 0), (0.4, 10)], "B2", angulo=90.0)]
    a = AssociadorGeometrico(barras, cfg)
    textos = [_texto(0.0, 5.0, rot=90.0), _texto(0.4, 5.2, rot=90.0)]
    res = a.associar_grupo(agrupar(textos, cfg)[0], textos)
    escolhidas = {r.indice_barra for r in res}
    assert len(escolhidas) == 2, "a mesma barra foi usada duas vezes"


def test_texto_sem_barra_no_raio_devolve_associacao_vazia(cfg):
    cfg.calibrar(0.2)
    a = AssociadorGeometrico([_barra([(500, 500), (500, 510)])], cfg)
    textos = [_texto(0.0, 0.0)]
    res = a.associar_grupo(agrupar(textos, cfg)[0], textos)
    assert res[0].indice_barra is None


def test_linha_tracejada_de_distribuicao_nao_vira_barra(cfg):
    cfg.calibrar(0.2)
    a = AssociadorGeometrico([_barra([(0, 0), (0, 10)], distribuicao=True)], cfg)
    assert a.barras == []


def test_linha_curta_demais_nao_vira_barra(cfg):
    """Dobras soltas e marcas nao podem ser confundidas com o corpo da barra."""
    cfg.calibrar(0.2)                       # minimo = 2 x 0.2 = 0.4
    a = AssociadorGeometrico([_barra([(0, 0), (0, 0.1)])], cfg)
    assert a.barras == []


def test_ponto_representativo_fica_sobre_a_barra_em_L(cfg):
    """O centroide de uma barra em L cai fora dela; o ponto medio nao."""
    barra = _barra([(0, 0), (10, 0), (10, 10)])
    x, y = ponto_representativo(barra, _texto(0, 0))
    assert barra.geometria.distance(LineString([(x, y), (x, y + 1e-9)])) < 1e-6


def test_ponto_representativo_cai_no_texto_sem_barra(cfg):
    assert ponto_representativo(None, _texto(7.0, 3.0)) == (7.0, 3.0)


# =============================================================================
# Ponta a ponta na prancha real
# =============================================================================

@pytest.mark.slow
def test_pipeline_completo_na_prancha_de_exemplo(cfg, dxf_referencia, tmp_path):
    from extrator.exportacao import Exportador
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    # Este teste mede a LEITURA do desenho, entao fixa a fonte da
    # quantidade no texto. O modo padrao (gabarito_rateado) e coberto
    # pelos testes de test_tabela_mestre.py.
    cfg.calculo.fonte_quantidade = "desenho"
    resultado = processar(dxf_referencia, cfg)
    e = resultado.estatisticas

    # --- leitura ---------------------------------------------------
    assert e.arquivos_lidos == 1
    assert e.textos_nao_interpretados == 0, \
        "todo texto de armadura precisa ser interpretado ou ignorado"
    # O desenho de referencia tem so o trecho: 6 textos de armadura, e
    # 3 deles rendem duas posicoes cada (armadura empilhada).
    assert e.textos_interpretados == 6
    assert e.posicoes_geradas == 9

    # --- o contorno do trecho sai da propria layer "TRECHO A" --------
    assert e.areas_encontradas == 1
    assert resultado.areas[0].nome == "TRECHO A"

    # --- pesos por bitola --------------------------------------------
    por_bitola = {}
    for p in resultado.posicoes:
        por_bitola.setdefault(p.bitola_mm, 0.0)
        por_bitola[p.bitola_mm] += p.comprimento_total_m
    assert por_bitola[8.0] == pytest.approx(244.75, abs=0.5)
    assert por_bitola[10.0] == pytest.approx(24.10, abs=0.5)

    liquido = sum(p.peso_liquido_kg for p in resultado.posicoes)
    com_perda = sum(p.peso_com_perda_kg for p in resultado.posicoes)
    assert liquido == pytest.approx(111.55, rel=0.001)
    assert com_perda == pytest.approx(liquido * 1.10, rel=1e-9)

    # --- soma das areas TEM que fechar com o total geral -------------
    soma_areas = sum(p.peso_com_perda_kg for p in resultado.posicoes)
    assert soma_areas == pytest.approx(com_perda, rel=1e-12)

    # --- Excel gerado com todas as abas exigidas ---------------------
    destino = Exportador(cfg).exportar(resultado, tmp_path / "saida.xlsx")
    assert destino.is_file()

    import openpyxl
    wb = openpyxl.load_workbook(destino)
    assert "RESUMO GERAL" in wb.sheetnames
    assert "VERIFICACAO" in wb.sheetnames
    assert "INCONSISTENCIAS" in wb.sheetnames
    assert any("TRECHO A" in a for a in wb.sheetnames)
    # cabecalho congelado e autofiltro
    assert wb["RESUMO GERAL"].freeze_panes == "B5"
    assert wb["RESUMO GERAL"].auto_filter.ref is not None


@pytest.mark.slow
def test_criterio_proporcional_nao_altera_o_total_geral(cfg, dxf_referencia, tmp_path):
    """Mudar o criterio de divisa redistribui entre areas, nunca cria aco."""
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    cfg.calculo.fonte_quantidade = "desenho"
    cfg.areas.criterio_divisa = "centroide"
    total_centroide = sum(p.peso_liquido_kg
                          for p in processar(dxf_referencia, cfg).posicoes)
    try:
        cfg.areas.criterio_divisa = "proporcional"
        total_prop = sum(p.peso_liquido_kg
                         for p in processar(dxf_referencia, cfg).posicoes)
    finally:
        cfg.areas.criterio_divisa = "centroide"
    assert total_prop == pytest.approx(total_centroide, rel=1e-9)


# =============================================================================
# Nome das abas com o prefixo do projeto
# =============================================================================

def test_nome_da_aba_recebe_o_prefixo_do_projeto():
    from extrator.exportacao import _nome_aba
    usados: set[str] = set()
    assert _nome_aba("AREA-01", usados, "Edificio X") == "Edificio X - AREA-01"


def test_prefixo_do_projeto_e_identico_em_todas_as_abas():
    """O mesmo projeto nao pode aparecer escrito de tres jeitos diferentes
    na barra de abas."""
    from extrator.exportacao import _nome_aba, _prefixo_comum
    sufixos = ["RESUMO GERAL", "VERIFICACAO", "INCONSISTENCIAS",
               "AREA-01", "FORA_DO_ESCOPO", "SEM_AREA"]
    prefixo = _prefixo_comum("Prancha Exemplo Aco", sufixos)
    usados: set[str] = set()
    nomes = [_nome_aba(s, usados, prefixo) for s in sufixos]
    assert all(len(n) <= 31 for n in nomes), nomes
    assert all(n.startswith(prefixo + " - ") for n in nomes), nomes


def test_projeto_muito_longo_faz_abreviar_o_sufixo():
    from extrator.exportacao import _nome_aba, _prefixo_comum
    sufixos = ["INCONSISTENCIAS"]
    prefixo = _prefixo_comum("Residencial Jardim das Palmeiras Bloco A", sufixos)
    nome = _nome_aba("INCONSISTENCIAS", set(), prefixo)
    assert len(nome) <= 31
    assert "INCONSIST" in nome


def test_nomes_de_aba_nunca_se_repetem():
    from extrator.exportacao import _nome_aba
    usados: set[str] = set()
    a = _nome_aba("AREA-01", usados, "Obra")
    b = _nome_aba("AREA-01", usados, "Obra")
    assert a != b and len(b) <= 31


def test_caracteres_proibidos_pelo_excel_sao_trocados():
    from extrator.exportacao import _nome_aba
    nome = _nome_aba("AREA/01", set(), "Obra:X")
    assert not set(nome) & set("[]:*?/\\")


@pytest.mark.slow
def test_abas_da_prancha_real_saem_com_o_prefixo(cfg, dxf_referencia, tmp_path):
    from extrator.exportacao import Exportador
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar
    import openpyxl

    configurar_log(cfg.log, tmp_path)
    cfg.projeto = "Obra Teste"
    try:
        resultado = processar(dxf_referencia, cfg)
        destino = Exportador(cfg).exportar(resultado, tmp_path / "s.xlsx")
    finally:
        cfg.projeto = ""
    abas = openpyxl.load_workbook(destino).sheetnames
    assert all(a.startswith("Obra Teste - ") for a in abas), abas
    assert all(len(a) <= 31 for a in abas)


@pytest.mark.slow
def test_barra_cortada_nao_gera_alerta_na_prancha_real(cfg, dxf_referencia, tmp_path):
    """No desenho de referencia o projetista ja apagou tudo que nao
    pertence ao trecho, entao NENHUMA barra cruza o contorno.

    O que este teste protege e justamente isso: fluxo limpo nao pode
    gerar posicao fora do escopo nem alerta nenhum. O corte em si tem
    cobertura sintetica em test_trechos.py, onde da para desenhar a
    barra exatamente em cima da divisa."""
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    cfg.areas.criterio_divisa = "fora_do_escopo"
    resultado = processar(dxf_referencia, cfg)

    cortadas = [p for p in resultado.posicoes
                if p.area == cfg.areas.nome_fora_escopo]
    assert not cortadas, f"nada devia estar fora do escopo: {cortadas[:3]}"
    assert all(p.area == "TRECHO A" for p in resultado.posicoes)

    erros = [i for i in resultado.inconsistencias
             if i.severidade.value == "ERRO"]
    assert not erros, f"desenho limpo gerou erro: {erros[:3]}"


@pytest.mark.slow
def test_criterio_fora_do_escopo_preserva_o_total_geral(cfg, dxf_referencia, tmp_path):
    """Tirar a barra do trecho nao pode tirar o aco do quantitativo."""
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    cfg.areas.criterio_divisa = "centroide"
    total_centroide = sum(p.peso_liquido_kg
                          for p in processar(dxf_referencia, cfg).posicoes)
    cfg.areas.criterio_divisa = "fora_do_escopo"
    total_escopo = sum(p.peso_liquido_kg
                       for p in processar(dxf_referencia, cfg).posicoes)
    assert total_escopo == pytest.approx(total_centroide, rel=1e-9)


# =============================================================================
# MODO ENXUTO SOBRE A PRANCHA DE REFERENCIA
# =============================================================================

@pytest.mark.slow
def test_conferencia_reconhece_todas_as_barras_da_prancha(cfg, dxf_referencia,
                                                          tmp_path):
    """As 9 posicoes desenhadas constam na tabela do projeto, com o mesmo
    comprimento unitario. Se alguma parar de bater, a leitura mudou."""
    import openpyxl

    from extrator.exportacao import Exportador
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    cfg.saida.modo_enxuto = True
    cfg.projeto = "SEM NADA"
    resultado = processar(dxf_referencia, cfg)
    destino = Exportador(cfg).exportar(resultado, tmp_path / "enxuto.xlsx")

    wb = openpyxl.load_workbook(destino)
    # So o trecho e a conferencia - nada de resumo, comparacao ou avisos.
    assert sorted(wb.sheetnames) == ["SEM NADA - CONFERENCIA",
                                     "SEM NADA - TRECHO A"]

    ws = wb["SEM NADA - CONFERENCIA"]
    linhas = [r for r in ws.iter_rows(values_only=True)]
    situacoes = [r[7] for r in linhas
                 if r[7] and r[7] not in ("Situacao",)]
    assert len(situacoes) == 9
    assert set(situacoes) == {"CONFERE"}


@pytest.mark.slow
def test_conferencia_nao_reprova_por_quantidade(cfg, dxf_referencia, tmp_path):
    """O desenho e um recorte: a tabela do projeto descreve o pavimento
    inteiro, entao a quantidade diverge de proposito. Isso NAO e problema
    - so a identidade da barra (posicao, bitola, comprimento) reprova."""
    from extrator.exportacao import Exportador
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    cfg.saida.modo_enxuto = True
    resultado = processar(dxf_referencia, cfg)

    exp = Exportador(cfg)
    from extrator.calculo import montar_dataframe
    df = montar_dataframe(resultado.posicoes, cfg)
    linhas = exp._linhas_conferencia(df, resultado)

    # Ha divergencia de quantidade de verdade nesta prancha...
    assert any(r["qtd_tabela"] and r["qtd_desenho"] != r["qtd_tabela"]
               for r in linhas)
    # ...e mesmo assim nenhuma barra e reprovada.
    assert {r["situacao"] for r in linhas} == {"CONFERE"}


@pytest.mark.slow
def test_conferencia_acusa_posicao_que_nao_esta_na_tabela(cfg, dxf_referencia,
                                                          tmp_path):
    """Barra desenhada que o projeto nao declara e o erro que a aba existe
    para pegar - o aco entra no corte sem respaldo da tabela."""
    from extrator.exportacao import Exportador
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    cfg.saida.modo_enxuto = True
    resultado = processar(dxf_referencia, cfg)

    # Some com a tabela mestre de uma posicao so.
    alvo = resultado.posicoes[0].posicao
    resultado.tabela_mestre = [l for l in resultado.tabela_mestre
                               if str(l.posicao) != str(alvo)]

    exp = Exportador(cfg)
    from extrator.calculo import montar_dataframe
    linhas = exp._linhas_conferencia(
        montar_dataframe(resultado.posicoes, cfg), resultado)
    assert "NAO CONSTA NA TABELA" in {r["situacao"] for r in linhas}
