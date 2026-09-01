"""
Testes da leitura da TABELA MESTRE desenhada na prancha (o gabarito).

O risco desta funcionalidade e ler um numero errado em silencio: a tabela
e uma grade de textos soltos, e uma celula intrusa faria o peso sair
completamente fora. Por isso cada linha e validada pela fisica
(peso ~= comprimento x peso linear) e os testes cobrem justamente os
casos em que a grade "suja".
"""
import pytest

from extrator.modelos import TextoDXF
from extrator.parser_texto import ParserArmadura
from extrator.tabela_mestre import LeitorTabelaMestre, totais_por_bitola


def _celula(x, y, texto, altura=0.2):
    return TextoDXF(handle="H", conteudo=texto, conteudo_bruto=texto,
                    x=x, y=y, altura=altura, rotacao=0.0,
                    layer="TEXTO_TABELAS", prancha="P")


def _linha_tabela(pernas, qtd=10, bitola="8", peso_linear=0.395):
    """Monta uma linha de tabela como o Eberick desenha.

    Ordem das celulas:  pos | bitola | qtd | <pernas> | comp.unit |
                        comp.total | peso
    """
    comp_unit = sum(float(x) for x in pernas)
    comp_total = comp_unit * qtd
    peso = comp_total / 100 * peso_linear
    linha = [_celula(270.0, 10.0, "1"), _celula(270.6, 10.0, f"%%c{bitola}"),
             _celula(271.5, 10.0, str(qtd))]
    x = 272.0
    for v in pernas:
        linha.append(_celula(x, 10.0, str(v)))
        x += 0.7
    linha.append(_celula(x, 10.0, f"{comp_unit:g}"))          # Comp. unitario
    linha.append(_celula(275.8, 10.0, f"{comp_total:.0f}"))   # Comp. total
    linha.append(_celula(277.0, 10.0, f"{peso:.2f}"))         # Peso
    return linha


@pytest.fixture
def leitor(cfg):
    return LeitorTabelaMestre(cfg)


@pytest.fixture
def normalizar(cfg):
    return ParserArmadura(cfg).normalizar


# =============================================================================
# Leitura basica
# =============================================================================

def test_le_uma_linha_simples(leitor, normalizar):
    """3 | %%c6.3 | 50 | ... | 111 | 5550 | 13.6
    50 barras de 111 cm = 5550 cm; 55,5 m x 0,245 = 13,6 kg."""
    linha = [_celula(x, 10.0, t) for x, t in [
        (270.1, "3"), (270.6, "%%c6.3"), (271.6, "50"), (272.5, "9"),
        (273.2, "93"), (274.1, "9"), (274.9, "111"), (275.9, "5550"),
        (277.0, "13.6")]]
    lidas, _ = leitor.ler(linha, normalizar)
    assert len(lidas) == 1
    l = lidas[0]
    assert l.posicao == "3"
    assert l.bitola_mm == 6.3
    assert l.quantidade == 50
    assert l.comprimento_total_cm == 5550
    assert l.peso_kg == pytest.approx(13.6)
    assert l.comprimento_total_m == pytest.approx(55.5)


def test_le_linha_com_nome_de_elemento_antes(leitor, normalizar):
    """A primeira celula pode ser o nome do elemento, nao a posicao."""
    linha = [_celula(x, 10.0, t) for x, t in [
        (266.8, "Armadura longitudinal"), (270.1, "1"), (270.6, "%%c10"),
        (271.5, "159"), (272.9, "VAR."), (274.8, "VAR."),
        (275.6, "142941"), (276.8, "880.8")]]
    lidas, _ = leitor.ler(linha, normalizar)
    assert len(lidas) == 1
    assert lidas[0].posicao == "1" and lidas[0].quantidade == 159


def test_celulas_VAR_no_meio_nao_atrapalham(leitor, normalizar):
    linha = [_celula(x, 10.0, t) for x, t in [
        (270.0, "9"), (270.6, "%%c6.3"), (271.5, "155"), (272.3, "19"),
        (272.9, "VAR."), (274.0, "19"), (274.8, "VAR."),
        (275.8, "18290"), (277.0, "44.8")]]
    lidas, _ = leitor.ler(linha, normalizar)
    assert len(lidas) == 1 and lidas[0].peso_kg == pytest.approx(44.8)


# =============================================================================
# Robustez - o ponto critico
# =============================================================================

def test_celula_intrusa_no_fim_e_aparada(leitor, normalizar):
    """O bloco de totais da direita cai no MESMO Y de uma linha de posicao.

    Sem aparar, o peso lido seria 19960 em vez de 18.5 - erro de 1000x
    passando despercebido.
    """
    linha = [_celula(x, 10.0, t) for x, t in [
        (270.0, "10"), (270.6, "%%c6.3"), (271.6, "23"), (272.3, "19"),
        (272.9, "VAR."), (274.0, "19"), (274.8, "VAR."),
        (275.9, "7544"), (277.0, "18.5"),
        # intrusos do bloco de totais, mesmo Y:
        (284.9, "2529.2"), (286.4, "3991"), (287.3, "19960")]]
    lidas, _ = leitor.ler(linha, normalizar)
    assert len(lidas) == 1
    assert lidas[0].peso_kg == pytest.approx(18.5), "leu o total geral como peso"
    assert lidas[0].comprimento_total_cm == pytest.approx(7544)


def test_linha_que_nao_fecha_a_conta_e_descartada(leitor, normalizar):
    """Peso incompativel com o comprimento: melhor perder a linha do que
    contaminar a comparacao."""
    linha = [_celula(x, 10.0, t) for x, t in [
        (270.0, "99"), (270.6, "%%c10"), (271.5, "10"),
        (275.6, "1000"), (276.8, "999999")]]
    lidas, avisos = leitor.ler(linha, normalizar)
    assert lidas == []
    assert any("nao fecharam a conta" in a for a in avisos)


def test_linha_de_cabecalho_e_ignorada(leitor, normalizar):
    linha = [_celula(x, 10.0, t) for x, t in [
        (267.4, "Elemento"), (269.8, "Pos."), (270.5, "Diam."),
        (271.5, "Q."), (284.4, "Comp. total"), (286.4, "Peso")]]
    lidas, _ = leitor.ler(linha, normalizar)
    assert lidas == []


def test_bloco_de_totais_sozinho_nao_vira_posicao(leitor, normalizar):
    linha = [_celula(x, 10.0, t) for x, t in [
        (283.3, "%%c6.3"), (284.9, "3436.7"), (286.6, "842")]]
    lidas, _ = leitor.ler(linha, normalizar)
    assert lidas == []


def test_bitola_fora_da_tabela_nbr_e_recusada(leitor, normalizar):
    linha = [_celula(x, 10.0, t) for x, t in [
        (270.0, "1"), (270.6, "%%c7.7"), (271.5, "10"),
        (275.6, "1000"), (276.8, "5.0")]]
    lidas, _ = leitor.ler(linha, normalizar)
    assert lidas == []


def test_sem_textos_devolve_aviso(leitor, normalizar):
    lidas, avisos = leitor.ler([], normalizar)
    assert lidas == [] and avisos


# =============================================================================
# Agregacao
# =============================================================================

def test_totais_por_bitola_somam(leitor, normalizar):
    celulas = []
    for i, (y, pos, bit, qtd, comp, peso) in enumerate([
            (10.0, "1", "%%c10", "100", "40000", "246.7"),
            (9.0, "2", "%%c10", "50", "20000", "123.4"),
            (8.0, "3", "%%c8", "10", "5000", "19.8")]):
        celulas += [_celula(270.0, y, pos), _celula(270.6, y, bit),
                    _celula(271.5, y, qtd), _celula(275.6, y, comp),
                    _celula(276.8, y, peso)]
    lidas, _ = leitor.ler(celulas, normalizar)
    assert len(lidas) == 3
    tot = totais_por_bitola(lidas)
    assert tot[10.0]["posicoes"] == 2
    assert tot[10.0]["peso_kg"] == pytest.approx(370.1)
    assert tot[10.0]["comprimento_m"] == pytest.approx(600.0)
    assert tot[8.0]["posicoes"] == 1


# =============================================================================
# Contra a prancha real - o gabarito precisa bater com ele mesmo
# =============================================================================

@pytest.mark.slow
def test_tabela_mestre_da_prancha_real_bate_com_os_totais_impressos(
        cfg, dxf_referencia, tmp_path):
    """A propria tabela imprime os totais por bitola no canto. O que o
    parser soma linha a linha tem de reproduzir esses numeros."""
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    resultado = processar(dxf_referencia, cfg)
    assert resultado.tabela_mestre, "nao leu a tabela mestre da prancha"

    assert len(resultado.tabela_mestre) == 268

    tot = totais_por_bitola(resultado.tabela_mestre)
    # Pesos impressos no bloco "Resumo Aco" do proprio desenho. Somar as
    # 268 linhas uma a uma tem de reproduzir exatamente estes numeros -
    # e a unica conferencia da leitura que nao depende de mim.
    impresso = {6.3: 160.8, 8.0: 2298.1, 10.0: 1662.5,
                12.5: 3520.4, 16.0: 1238.2, 20.0: 39.5}
    for bitola, peso_kg in impresso.items():
        assert tot[bitola]["peso_kg"] == pytest.approx(peso_kg, abs=0.1), bitola

    # "Total: 8919.5", tambem impresso na prancha.
    total_lido = sum(v["peso_kg"] for v in tot.values())
    assert total_lido == pytest.approx(8919.5, abs=0.5)


@pytest.mark.slow
def test_tabela_mestre_nao_entra_no_quantitativo(cfg, dxf_referencia, tmp_path):
    """O gabarito e so referencia: se entrasse no total, o aco seria
    contado duas vezes."""
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    cfg.calculo.fonte_quantidade = "desenho"
    resultado = processar(dxf_referencia, cfg)
    extraido = sum(p.peso_liquido_kg for p in resultado.posicoes)
    mestre = sum(l.peso_kg for l in resultado.tabela_mestre)
    assert extraido == pytest.approx(111.55, rel=0.001)
    assert mestre == pytest.approx(8919.5, abs=0.5)
    assert extraido < mestre


@pytest.mark.slow
def test_aba_comparacao_traz_o_percentual_de_acerto(cfg, dxf_referencia, tmp_path):
    import openpyxl

    from extrator.exportacao import Exportador
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    resultado = processar(dxf_referencia, cfg)
    destino = Exportador(cfg).exportar(resultado, tmp_path / "c.xlsx")
    wb = openpyxl.load_workbook(destino)
    aba = [a for a in wb.sheetnames if "COMPARACAO" in a.upper()][0]
    texto = " ".join(str(v) for linha in wb[aba].iter_rows(values_only=True)
                     for v in linha if v is not None)
    assert "% DE ACERTO" in texto
    assert "Tabela mestre (kg)" in texto
    assert "Gabarito lido da prancha" in texto


# =============================================================================
# Rateio do gabarito pelos trechos (modo padrao)
# =============================================================================

@pytest.mark.slow
def test_rateio_faz_o_total_bater_com_o_gabarito(cfg, dxf_referencia, tmp_path):
    """O rateio da a quantidade EXATA da tabela as posicoes desenhadas -
    e so a elas.

    O desenho de referencia e um trecho: tem 9 posicoes das 268 que a
    tabela descreve. Ratear o pavimento inteiro em cima dessas 9 seria o
    pior erro possivel (uma barra desenhada virando centenas), entao o
    que se exige aqui e o contrario do "acerto de 100%": cada posicao
    desenhada recebe a quantidade que a tabela da PARA ELA, e as 259 que
    nao foram desenhadas ficam de fora, reportadas."""
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    cfg.calculo.fonte_quantidade = "gabarito_rateado"
    res = processar(dxf_referencia, cfg)

    gab = {(l.posicao, round(l.bitola_mm, 3)): l for l in res.tabela_mestre}
    assert res.posicoes
    for p in res.posicoes:
        linha = gab.get((p.posicao, round(p.bitola_mm, 3)))
        assert linha is not None, f"N{p.posicao} nao esta na tabela"
        # trecho unico: a posicao leva a quantidade inteira da tabela
        assert p.quantidade == pytest.approx(linha.quantidade, rel=0.001), \
            f"N{p.posicao}"

    # e o peso NUNCA pode passar do que a tabela declara
    extraido = sum(p.peso_liquido_kg for p in res.posicoes)
    mestre = sum(l.peso_kg for l in res.tabela_mestre)
    assert extraido < mestre


@pytest.mark.slow
def test_rateio_guarda_a_quantidade_lida_do_desenho(cfg, dxf_referencia, tmp_path):
    """A leitura do desenho nao se perde: fica na coluna de conferencia."""
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    cfg.calculo.fonte_quantidade = "gabarito_rateado"
    res = processar(dxf_referencia, cfg)

    do_gabarito = [p for p in res.posicoes if p.origem_quantidade == "gabarito"]
    assert do_gabarito, "nenhuma posicao recebeu quantidade do gabarito"
    localizadas = [p for p in do_gabarito if p.quantidade_desenho > 0]
    assert localizadas, "quantidade do desenho nao foi preservada"
    # A quantidade rateada quase sempre difere da lida - e justamente o
    # que a coluna de conferencia serve para mostrar.
    assert any(abs(p.quantidade - p.quantidade_desenho) > 0.5 for p in localizadas)


@pytest.mark.slow
def test_posicao_do_gabarito_sem_local_nao_some(cfg, dxf_referencia, tmp_path):
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    cfg.calculo.fonte_quantidade = "gabarito_rateado"
    res = processar(dxf_referencia, cfg)

    sem_local = [p for p in res.posicoes
                 if p.area == cfg.areas.nome_sem_localizacao]
    if sem_local:
        # se existe, o peso dela tem de estar no total
        assert all(p.peso_liquido_kg > 0 for p in sem_local)
        tipos = [i.tipo.name for i in res.inconsistencias]
        assert "POSICAO_SEM_LOCALIZACAO" in tipos


def test_modo_desenho_nao_usa_o_gabarito(cfg):
    """Com fonte_quantidade='desenho' o gabarito e so referencia."""
    from extrator.rateio import ratear_por_gabarito
    from extrator.modelos import ResultadoProcessamento
    res = ResultadoProcessamento()
    # sem gabarito, a lista volta intacta
    assert ratear_por_gabarito([], [], cfg, "P", res) == []


# =============================================================================
# Formato da barra (dobras) - corte e dobra
# =============================================================================

@pytest.mark.parametrize("meio,esperado", [
    # (celulas do meio da tabela, (dobra1, reta, dobra2))
    (["400"],                (None, 400.0, None)),   # barra reta pura
    (["19", "391"],          (19.0, 391.0, None)),   # gancho na entrada
    (["391", "19"],          (None, 391.0, 19.0)),   # gancho na saida
    (["9", "93", "9"],       (9.0, 93.0, 9.0)),      # dois ganchos
    (["14", "255", "14"],    (14.0, 255.0, 14.0)),
])
def test_decomposicao_em_dobra_reta_dobra(leitor, normalizar, meio, esperado):
    """A RETA e a perna maior (corpo da barra); a ORDEM diz de que lado
    esta o gancho. Sem isso, "19 + 391" e "391 + 19" viram a mesma peca.

    Atencao ao montar a linha: na tabela do projeto, DEPOIS das pernas vem
    ainda a coluna "Comp. (cm)" com o comprimento unitario. A linha real de
    uma barra com dois ganchos tem quatro celulas no meio:
        9 | 93 | 9 | 111
        as tres primeiras sao as pernas, a ultima e o comprimento unitario.
    """
    lidas, _ = leitor.ler(_linha_tabela(meio, qtd=10, bitola="8",
                                        peso_linear=0.395), normalizar)
    assert len(lidas) == 1, f"nao leu {meio}"
    assert lidas[0].dobra_reta_dobra == esperado
    assert lidas[0].comprimento_unitario_cm == sum(float(x) for x in meio)


def test_formato_variavel_nao_tem_dobras(leitor, normalizar):
    linha = [_celula(x, 10.0, t) for x, t in [
        (270.0, "8"), (270.6, "%%c6.3"), (271.8, "7"), (272.3, "19"),
        (272.9, "VAR."), (274.0, "19"), (274.8, "VAR."),
        (275.9, "2219"), (277.1, "5.4")]]
    lidas, _ = leitor.ler(linha, normalizar)
    assert len(lidas) == 1
    l = lidas[0]
    assert l.variavel is True
    assert l.pernas_cm == ()
    assert l.chave_formato == (), "barra variavel nao pode casar com outra"


def test_chave_de_formato_separa_pecas_diferentes(leitor, normalizar):
    """Duas barras de 111 cm com dobras diferentes NAO sao a mesma peca."""
    def monta(pernas):
        linha = _linha_tabela(pernas, qtd=10, bitola="6.3", peso_linear=0.245)
        return leitor.ler(linha, normalizar)[0][0]

    a = monta([9, 93, 9])      # 111 cm
    b = monta([20, 71, 20])    # 111 cm, formato diferente
    assert a.comprimento_unitario_cm == b.comprimento_unitario_cm == 111
    assert a.chave_formato != b.chave_formato, \
        "pecas de formatos diferentes nao podem ser intercambiaveis"


@pytest.mark.slow
def test_dobras_chegam_na_planilha(cfg, dxf_referencia, tmp_path):
    """O caminho completo: tabela do projeto -> rateio -> Excel."""
    import openpyxl

    from extrator.exportacao import Exportador
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    res = processar(dxf_referencia, cfg)
    destino = Exportador(cfg).exportar(res, tmp_path / "d.xlsx")

    wb = openpyxl.load_workbook(destino)
    ws = wb[[n for n in wb.sheetnames if "RESUMO" in n.upper()][0]]
    linhas = list(ws.iter_rows(values_only=True))
    cab = next(r for r in linhas if r and "Posicao" in [str(x) for x in r])
    nomes = [str(x) for x in cab]
    for coluna in ("Dobra 1 (cm)", "Reta (cm)", "Dobra 2 (cm)"):
        assert coluna in nomes, f"falta a coluna {coluna}"

    # pelo menos uma barra com dobra preenchida
    i_d1 = nomes.index("Dobra 1 (cm)")
    i_reta = nomes.index("Reta (cm)")
    com_dobra = [r for r in linhas
                 if len(r) > i_d1 and isinstance(r[i_d1], (int, float))]
    assert com_dobra, "nenhuma barra saiu com dobra na planilha"
    # e a conta fecha: dobra + reta <= comprimento unitario
    i_comp = nomes.index("Comprimento unitario (cm)")
    r = com_dobra[0]
    assert r[i_d1] + r[i_reta] <= r[i_comp] + 0.01


# =============================================================================
# Dobras vindas da tabela quando a quantidade vem do desenho (recorte)
# =============================================================================

def test_formato_vem_da_tabela_mesmo_com_quantidade_do_desenho(cfg):
    """Num recorte a quantidade vem do desenho, mas o FORMATO da barra e
    propriedade da posicao - tem de vir da tabela do projeto, senao a
    planilha nao serve para corte e dobra."""
    from extrator.modelos import PosicaoArmadura, ResultadoProcessamento
    from extrator.rateio import aplicar_formato_do_gabarito
    from extrator.tabela_mestre import LinhaMestre

    linha = LinhaMestre(posicao="6", bitola_mm=8.0, quantidade=80,
                        comprimento_total_cm=22800, peso_kg=90.1)
    linha.pernas_cm = (14.0, 257.0, 14.0)
    linha.comprimento_unitario_cm = 285.0

    p = PosicaoArmadura(id="x", prancha="P", posicao="6", quantidade=1,
                        bitola_mm=8.0, comprimento_unit_cm=285.0,
                        espacamento_cm=10.0, texto_origem="N6-Ø8c/10 C=285")
    res = ResultadoProcessamento()
    aplicar_formato_do_gabarito([p], [linha], cfg, "P", res)

    assert (p.dobra1_cm, p.reta_cm, p.dobra2_cm) == (14.0, 257.0, 14.0)
    assert p.quantidade == 1, "a quantidade do desenho nao pode ser alterada"


def test_formato_nao_e_aplicado_se_o_comprimento_diverge(cfg):
    """Comprimento diferente significa que texto e tabela falam de barras
    diferentes. Forcar o formato ali seria inventar dado."""
    from extrator.modelos import PosicaoArmadura, ResultadoProcessamento
    from extrator.rateio import aplicar_formato_do_gabarito
    from extrator.tabela_mestre import LinhaMestre

    linha = LinhaMestre(posicao="6", bitola_mm=8.0, quantidade=80,
                        comprimento_total_cm=22800, peso_kg=90.1)
    linha.pernas_cm = (14.0, 257.0, 14.0)
    linha.comprimento_unitario_cm = 285.0

    p = PosicaoArmadura(id="x", prancha="P", posicao="6", quantidade=1,
                        bitola_mm=8.0, comprimento_unit_cm=999.0,
                        espacamento_cm=None, texto_origem="N6-Ø8 C=999")
    res = ResultadoProcessamento()
    aplicar_formato_do_gabarito([p], [linha], cfg, "P", res)

    assert p.dobra1_cm is None and p.reta_cm is None
    assert "COMPRIMENTO_DIVERGE_DO_GABARITO" in [
        i.tipo.name for i in res.inconsistencias]


@pytest.mark.slow
def test_recorte_real_sai_com_as_dobras(cfg, tmp_path, dxf_referencia):
    """Ponta a ponta no recorte que o usuario mandou."""
    from extrator.log_config import configurar_log
    from extrator.pipeline import processar

    configurar_log(cfg.log, tmp_path)
    res = processar(dxf_referencia, cfg)

    por_pos = {p.posicao: p for p in res.posicoes}
    # valores conferidos na tabela do proprio desenho
    assert (por_pos["6"].dobra1_cm, por_pos["6"].reta_cm,
            por_pos["6"].dobra2_cm) == (14.0, 257.0, 14.0)
    assert (por_pos["13"].dobra1_cm, por_pos["13"].reta_cm,
            por_pos["13"].dobra2_cm) == (14.0, 533.0, 14.0)
    assert (por_pos["264"].dobra1_cm, por_pos["264"].reta_cm) == (19.0, 771.0)
    # e o trecho foi reconhecido apesar do canto solto
    assert [a.nome for a in res.areas] == ["TRECHO A"]


# =============================================================================
# A COLUNA MANDA NO FORMATO
# =============================================================================
# A tabela do Eberick tem colunas proprias para as dobras:
#
#     Elemento | Pos. | Diam. | Q. | Dob. | Reta | Dob. | Comp. | Total | CA-50
#
# Celula vazia NAO existe como texto no DXF: uma barra reta simplesmente
# nao tem a celula "Dob.". Contar celulas da esquerda para a direita, ou
# supor que "a maior perna e a reta", erra o lado do gancho. O X de cada
# celula - que e o mesmo de toda a coluna, porque o texto e alinhado -
# diz exatamente em que coluna o projetista escreveu cada numero.

# X reais medidos na tabela do projeto (SEM NADA.dxf):
X_POS, X_BIT, X_QTD = 257.56, 258.50, 259.37
X_DOB1, X_RETA, X_DOB2 = 260.14, 261.13, 261.96
X_COMP, X_TOTAL, X_PESO = 263.14, 264.34, 265.40

# ... e os X do cabecalho, que sao ligeiramente diferentes dos dados
# (o rotulo e mais largo que o numero).
CABECALHO = [(257.67, "Pos."), (258.79, "Diam."), (259.32, "Q."),
             (260.31, "Dob."), (261.20, "Reta"), (262.14, "Dob."),
             (263.39, "Comp."), (264.37, "Total"), (265.79, "CA-50")]


def _tabela(linhas_de_dados, com_cabecalho=True):
    """Monta a tabela com cabecalho, como ela existe no desenho."""
    celulas = []
    if com_cabecalho:
        celulas += [_celula(x, 115.49, t) for x, t in CABECALHO]
    for i, campos in enumerate(linhas_de_dados):
        y = 114.5 - i * 0.36
        for x, texto in campos:
            celulas.append(_celula(x, y, texto))
    return celulas


def _dados(pos, bitola, qtd, dob1, reta, dob2, peso_linear=0.395):
    """Uma linha de dados; dob1/dob2 None viram celula INEXISTENTE."""
    comp = (dob1 or 0) + reta + (dob2 or 0)
    total = comp * qtd
    campos = [(X_POS, str(pos)), (X_BIT, f"%%c{bitola:g}"), (X_QTD, str(qtd))]
    if dob1 is not None:
        campos.append((X_DOB1, f"{dob1:g}"))
    campos.append((X_RETA, f"{reta:g}"))
    if dob2 is not None:
        campos.append((X_DOB2, f"{dob2:g}"))
    campos += [(X_COMP, f"{comp:g}"), (X_TOTAL, f"{total:.0f}"),
               (X_PESO, f"{total / 100 * peso_linear:.1f}")]
    return campos


@pytest.mark.parametrize("dob1,reta,dob2", [
    (None, 230, None),   # N1 real: barra reta
    (14, 116, None),     # N2 real: gancho so na entrada
    (14, 1096, 14),      # N4 real: gancho dos dois lados
    (19, 291, None),     # N10 real
    (14, 533, 14),       # N13 real
])
def test_le_as_dobras_na_coluna_certa(leitor, normalizar, dob1, reta, dob2):
    textos = _tabela([_dados(1, 8, 10, dob1, reta, dob2)])
    lidas, _ = leitor.ler(textos, normalizar)
    assert len(lidas) == 1
    l = lidas[0]
    assert l.formato_por_coluna
    assert (l.dobra_1_cm, l.reta_cm, l.dobra_2_cm) == (dob1, reta, dob2)
    assert l.dobra_reta_dobra == (dob1, reta, dob2)
    assert l.comprimento_unitario_cm == (dob1 or 0) + reta + (dob2 or 0)
    assert l.formato_conferido


def test_gancho_maior_que_a_reta_nao_troca_de_coluna(leitor, normalizar):
    """A heuristica antiga chamava de "reta" a perna maior. A coluna nao
    se importa com o tamanho: 60 na coluna Dob. e dobra, e ponto."""
    textos = _tabela([_dados(1, 8, 10, 60, 40, None)])
    lidas, _ = leitor.ler(textos, normalizar)
    l = lidas[0]
    assert (l.dobra_1_cm, l.reta_cm, l.dobra_2_cm) == (60, 40, None)


def test_dobra_so_na_saida_nao_vira_dobra_na_entrada(leitor, normalizar):
    """Celula na coluna Dob. da DIREITA: o gancho esta na ponta de saida."""
    textos = _tabela([_dados(1, 8, 10, None, 400, 19)])
    lidas, _ = leitor.ler(textos, normalizar)
    l = lidas[0]
    assert (l.dobra_1_cm, l.reta_cm, l.dobra_2_cm) == (None, 400, 19)


def test_sem_cabecalho_avisa_que_o_formato_e_palpite(leitor, normalizar):
    textos = _tabela([_dados(1, 8, 10, 14, 286, None)], com_cabecalho=False)
    lidas, avisos = leitor.ler(textos, normalizar)
    assert len(lidas) == 1
    assert not lidas[0].formato_por_coluna
    assert any("cabecalho" in a for a in avisos)


def test_varias_linhas_com_formatos_diferentes(leitor, normalizar):
    """O que quebraria uma leitura por contagem de celulas: linhas
    vizinhas com numero DIFERENTE de celulas preenchidas."""
    textos = _tabela([
        _dados(1, 8, 23, None, 230, None),
        _dados(2, 8, 396, 14, 116, None),
        _dados(3, 8, 609, 14, 286, None),
        _dados(4, 8, 75, 14, 1096, 14),
        _dados(12, 8, 25, None, 560, None),
        _dados(13, 8, 25, 14, 533, 14),
    ])
    lidas, _ = leitor.ler(textos, normalizar)
    assert len(lidas) == 6
    achado = {l.posicao: l.dobra_reta_dobra for l in lidas}
    assert achado == {
        "1": (None, 230, None),
        "2": (14, 116, None),
        "3": (14, 286, None),
        "4": (14, 1096, 14),
        "12": (None, 560, None),
        "13": (14, 533, 14),
    }
    assert all(l.formato_conferido for l in lidas)
