"""
Testes dos trechos por layer, das secoes por sentido e da aba COMPARACAO.

A prancha de exemplo tem uma layer "AREA" so; estes testes geram um DXF
sintetico com duas layers "TRECHO 01" e "TRECHO 02" para exercitar o modo
recomendado (uma layer por trecho).
"""
import ezdxf
import pytest

from extrator.exportacao import Exportador
from extrator.log_config import configurar_log
from extrator.pipeline import processar


# =============================================================================
# DXF sintetico
# =============================================================================

def _prancha_dois_trechos(caminho, contorno_2_fechado=True):
    """Duas regioes lado a lado, cada uma na sua layer.

    TRECHO 01 (x 0..100) e desenhado com 4 LINEs soltas, como na prancha
    real do escritorio. TRECHO 02 (x 100..200) usa polilinha fechada.

    Dentro de cada trecho ha duas barras: uma horizontal e uma vertical,
    para exercitar a separacao por sentido.
    """
    doc = ezdxf.new("R2010", setup=True)
    msp = doc.modelspace()
    for nome in ("TRECHO 01", "TRECHO 02", "ARR_TXT", "ARR_BARRA"):
        doc.layers.add(nome)

    # --- contorno do TRECHO 01: 4 linhas soltas ----------------------
    cantos = [(0, 0), (100, 0), (100, 100), (0, 100)]
    for a, b in zip(cantos, cantos[1:] + cantos[:1]):
        msp.add_line(a, b, dxfattribs={"layer": "TRECHO 01"})

    # --- contorno do TRECHO 02: polilinha fechada --------------------
    pontos2 = [(100, 0), (200, 0), (200, 100), (100, 100)]
    if contorno_2_fechado:
        msp.add_lwpolyline(pontos2, close=True,
                           dxfattribs={"layer": "TRECHO 02"})
    else:
        # Contorno faltando um lado: nao deve virar trecho nenhum.
        for a, b in zip(pontos2, pontos2[1:]):
            msp.add_line(a, b, dxfattribs={"layer": "TRECHO 02"})

    # --- armadura: (x, y, dx, dy, texto, rotacao) --------------------
    barras = [
        # TRECHO 01 - horizontal e vertical
        (10, 50, 60, 0, "10N1-%%c10c/15 C=400", 0),
        (50, 10, 0, 60, "20N2-%%c12.5c/20 C=300", 90),
        # TRECHO 02 - horizontal e vertical
        (110, 50, 60, 0, "5N3-%%c8 C=200", 0),
        (150, 10, 0, 60, "8N4-%%c16 C=500", 90),
        # fora de qualquer trecho
        (300, 300, 40, 0, "3N5-%%c10 C=100", 0),
    ]
    for x, y, dx, dy, texto, rot in barras:
        msp.add_line((x, y), (x + dx, y + dy),
                     dxfattribs={"layer": "ARR_BARRA"})
        t = msp.add_text(texto, dxfattribs={
            "layer": "ARR_TXT", "height": 2.0, "rotation": rot})
        t.dxf.insert = (x + dx / 2, y + dy / 2)

    doc.saveas(caminho)
    return caminho


@pytest.fixture
def prancha_trechos(tmp_path):
    return _prancha_dois_trechos(tmp_path / "dois_trechos.dxf")


@pytest.fixture
def cfg_trechos(cfg):
    """Config apontada para as layers do DXF sintetico."""
    cfg.layers.texto_armadura = ["ARR_*"]
    cfg.layers.geometria_barra = ["ARR_*"]
    return cfg


# =============================================================================
# Nome do trecho vem da layer
# =============================================================================

def test_cada_layer_trecho_vira_um_trecho(cfg_trechos, prancha_trechos, tmp_path):
    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(prancha_trechos, cfg_trechos)
    nomes = sorted(a.nome for a in res.areas)
    assert nomes == ["TRECHO 01", "TRECHO 02"]
    assert all(a.origem_nome == "nome_layer" for a in res.areas)


def test_contorno_de_linhas_soltas_e_de_polilinha_funcionam_igual(
        cfg_trechos, prancha_trechos, tmp_path):
    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(prancha_trechos, cfg_trechos)
    por_nome = {a.nome: a for a in res.areas}
    assert por_nome["TRECHO 01"].montado_de_segmentos is True    # 4 LINEs
    assert por_nome["TRECHO 02"].montado_de_segmentos is False   # LWPOLYLINE


def test_armadura_cai_no_trecho_certo(cfg_trechos, prancha_trechos, tmp_path):
    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(prancha_trechos, cfg_trechos)
    por_area = {}
    for p in res.posicoes:
        por_area.setdefault(p.area, set()).add(p.posicao)
    assert por_area["TRECHO 01"] == {"1", "2"}
    assert por_area["TRECHO 02"] == {"3", "4"}
    assert por_area[cfg_trechos.areas.nome_sem_area] == {"5"}


def test_layers_de_trechos_vizinhos_nao_se_misturam(cfg_trechos, prancha_trechos,
                                                    tmp_path):
    """Os contornos encostam em x=100. Poligonizados juntos, os segmentos
    formariam um retangulo unico de 0 a 200."""
    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(prancha_trechos, cfg_trechos)
    for a in res.areas:
        minx, _, maxx, _ = a.poligono.bounds
        assert maxx - minx == pytest.approx(100.0, abs=1.0), \
            f"{a.nome} tem largura {maxx - minx}, esperado 100"


def test_trecho_nao_fechado_e_reportado(cfg_trechos, tmp_path):
    caminho = _prancha_dois_trechos(tmp_path / "aberto.dxf",
                                    contorno_2_fechado=False)
    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(caminho, cfg_trechos)
    assert [a.nome for a in res.areas] == ["TRECHO 01"]
    tipos = [i.tipo.name for i in res.inconsistencias]
    assert "AREA_NAO_FECHADA" in tipos


def test_prefixo_da_layer_pode_ser_removido(cfg_trechos, prancha_trechos, tmp_path):
    cfg_trechos.areas.limpar_prefixo_layer = True
    try:
        configurar_log(cfg_trechos.log, tmp_path)
        res = processar(prancha_trechos, cfg_trechos)
        assert sorted(a.nome for a in res.areas) == ["01", "02"]
    finally:
        cfg_trechos.areas.limpar_prefixo_layer = False


# =============================================================================
# Sentido da armadura
# =============================================================================

def test_sentido_e_classificado_pelo_angulo_da_barra(cfg_trechos, prancha_trechos,
                                                     tmp_path):
    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(prancha_trechos, cfg_trechos)
    por_posicao = {p.posicao: p.sentido for p in res.posicoes}
    assert por_posicao["1"] == cfg_trechos.sentido.nome_horizontal
    assert por_posicao["2"] == cfg_trechos.sentido.nome_vertical
    assert por_posicao["3"] == cfg_trechos.sentido.nome_horizontal
    assert por_posicao["4"] == cfg_trechos.sentido.nome_vertical


def test_cada_trecho_tem_uma_aba_so_com_os_dois_sentidos_dentro(
        cfg_trechos, prancha_trechos, tmp_path):
    """O pedido: UMA aba por trecho, com separacao interna por sentido -
    e nao uma aba por sentido."""
    import openpyxl

    configurar_log(cfg_trechos.log, tmp_path)
    cfg_trechos.projeto = "Obra"
    try:
        res = processar(prancha_trechos, cfg_trechos)
        destino = Exportador(cfg_trechos).exportar(res, tmp_path / "s.xlsx")
    finally:
        cfg_trechos.projeto = ""

    wb = openpyxl.load_workbook(destino)
    abas = wb.sheetnames
    assert "Obra - TRECHO 01" in abas
    assert "Obra - TRECHO 02" in abas
    # nenhuma aba dedicada a um sentido
    assert not [a for a in abas if "SENTIDO" in a.upper()]

    ws = wb["Obra - TRECHO 01"]
    conteudo = "\n".join(
        str(c.value) for row in ws.iter_rows(values_only=True)
        for c in [type("C", (), {"value": v})() for v in row] if c.value)
    assert cfg_trechos.sentido.nome_horizontal in conteudo
    assert cfg_trechos.sentido.nome_vertical in conteudo
    assert "TOTAL DO TRECHO" in conteudo
    assert "RESUMO POR SENTIDO" in conteudo


# =============================================================================
# Aba COMPARACAO FINAL
# =============================================================================

def test_eficacia_e_a_fracao_do_aco_dentro_de_trechos(cfg_trechos,
                                                      prancha_trechos, tmp_path):
    from extrator.calculo import eficacia_por_bitola, montar_dataframe

    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(prancha_trechos, cfg_trechos)
    df = montar_dataframe(res.posicoes, cfg_trechos)
    fora = [cfg_trechos.areas.nome_sem_area, cfg_trechos.areas.nome_fora_escopo]

    dentro = df[~df["Area"].isin(fora)]["Peso com perda (kg)"].sum()
    total = df["Peso com perda (kg)"].sum()
    esperado = dentro / total

    tab = eficacia_por_bitola(df, fora, "Peso com perda (kg)")
    assert (tab["Tabela mestre (kg)"].sum()
            == pytest.approx(float(total), rel=1e-9))
    assert (tab["Total nos trechos (kg)"].sum()
            == pytest.approx(float(dentro), rel=1e-9))
    assert 0 < esperado < 1        # a N5 esta fora, entao nao pode dar 100%


def test_eficacia_100_quando_tudo_esta_em_trechos(cfg_trechos, tmp_path):
    """Contorno cobrindo toda a armadura -> meta atingida."""
    from extrator.calculo import eficacia_por_bitola, montar_dataframe

    caminho = tmp_path / "tudo_dentro.dxf"
    doc = ezdxf.new("R2010", setup=True)
    msp = doc.modelspace()
    for nome in ("TRECHO 01", "ARR_TXT", "ARR_BARRA"):
        doc.layers.add(nome)
    msp.add_lwpolyline([(0, 0), (500, 0), (500, 500), (0, 500)], close=True,
                       dxfattribs={"layer": "TRECHO 01"})
    msp.add_line((100, 100), (200, 100), dxfattribs={"layer": "ARR_BARRA"})
    t = msp.add_text("10N1-%%c10c/15 C=400",
                     dxfattribs={"layer": "ARR_TXT", "height": 2.0})
    t.dxf.insert = (150, 100)
    doc.saveas(caminho)

    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(caminho, cfg_trechos)
    df = montar_dataframe(res.posicoes, cfg_trechos)
    fora = [cfg_trechos.areas.nome_sem_area, cfg_trechos.areas.nome_fora_escopo]
    tab = eficacia_por_bitola(df, fora, "Peso com perda (kg)")
    assert float(tab["Eficacia"].iloc[0]) == pytest.approx(1.0)
    assert float(tab["Fora (kg)"].iloc[0]) == pytest.approx(0.0)


def test_aba_comparacao_existe_e_traz_o_total(cfg_trechos, prancha_trechos,
                                              tmp_path):
    import openpyxl

    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(prancha_trechos, cfg_trechos)
    destino = Exportador(cfg_trechos).exportar(res, tmp_path / "c.xlsx")
    wb = openpyxl.load_workbook(destino)
    aba = [a for a in wb.sheetnames if "COMPARACAO" in a.upper()]
    assert aba, wb.sheetnames
    texto = " ".join(str(v) for linha in wb[aba[0]].iter_rows(values_only=True)
                     for v in linha if v is not None)
    assert "TOTAL" in texto
    # Bloco 1: acerto contra o gabarito. Bloco 2: cobertura dos trechos.
    assert "% DE ACERTO" in texto
    assert "Cobertura" in texto
    assert "PESO POR TRECHO" in texto


# =============================================================================
# Modo limpo - planilha enxuta SEM perder a lista de corte e dobra
# =============================================================================

def _abrir(cfg, prancha, tmp_path, nome="s.xlsx"):
    import openpyxl
    configurar_log(cfg.log, tmp_path)
    cfg.projeto = "Obra"
    res = processar(prancha, cfg)
    destino = Exportador(cfg).exportar(res, tmp_path / nome)
    return openpyxl.load_workbook(destino), res


def _texto_da_aba(ws) -> str:
    return "\n".join(str(v) for linha in ws.iter_rows(values_only=True)
                     for v in linha if v is not None)


def test_modo_limpo_gera_so_trechos_resumo_e_comparacao(cfg_trechos,
                                                        prancha_trechos, tmp_path):
    cfg_trechos.saida.modo_limpo = True
    wb, _ = _abrir(cfg_trechos, prancha_trechos, tmp_path)
    abas = wb.sheetnames
    assert not [a for a in abas if "VERIFICACAO" in a.upper()]
    assert not [a for a in abas if "INCONSIST" in a.upper()]
    assert [a for a in abas if "RESUMO GERAL" in a.upper()]
    assert [a for a in abas if "COMPARACAO" in a.upper()]
    assert "Obra - TRECHO 01" in abas and "Obra - TRECHO 02" in abas


def test_modo_limpo_MANTEM_comprimento_unitario_e_posicao(cfg_trechos,
                                                          prancha_trechos,
                                                          tmp_path):
    """E lista de corte e dobra: a medida exata de cada barra tem de estar
    la. Agregar por bitola apagaria justamente essa informacao."""
    cfg_trechos.saida.modo_limpo = True
    wb, _ = _abrir(cfg_trechos, prancha_trechos, tmp_path)
    texto = _texto_da_aba(wb["Obra - TRECHO 01"])
    assert "Posicao" in texto
    assert "Comprimento unitario (cm)" in texto
    assert "N1" in texto and "N2" in texto          # as posicoes
    assert "400" in texto and "300" in texto        # os comprimentos unitarios


def test_modo_limpo_tira_o_detalhamento_linha_a_linha(cfg_trechos,
                                                      prancha_trechos, tmp_path):
    cfg_trechos.saida.modo_limpo = True
    wb, _ = _abrir(cfg_trechos, prancha_trechos, tmp_path)
    for aba in wb.sheetnames:
        assert "DETALHAMENTO" not in _texto_da_aba(wb[aba]).upper(), aba


def test_modo_limpo_avisa_das_inconsistencias_no_resumo(cfg_trechos,
                                                        prancha_trechos, tmp_path):
    """Sem a aba INCONSISTENCIAS, o balanco nao pode desaparecer."""
    cfg_trechos.saida.modo_limpo = True
    wb, _ = _abrir(cfg_trechos, prancha_trechos, tmp_path)
    resumo = [a for a in wb.sheetnames if "RESUMO GERAL" in a.upper()][0]
    texto = _texto_da_aba(wb[resumo]).upper()
    assert ("ATENCAO" in texto) or ("NENHUM ERRO OU ALERTA" in texto)


def test_modo_limpo_mantem_as_secoes_por_sentido(cfg_trechos, prancha_trechos,
                                                 tmp_path):
    cfg_trechos.saida.modo_limpo = True
    wb, _ = _abrir(cfg_trechos, prancha_trechos, tmp_path)
    texto = _texto_da_aba(wb["Obra - TRECHO 01"])
    assert cfg_trechos.sentido.nome_horizontal in texto
    assert cfg_trechos.sentido.nome_vertical in texto
    assert "TOTAL DO TRECHO" in texto


def test_modo_limpo_nao_altera_nenhum_numero(cfg_trechos, prancha_trechos,
                                             tmp_path):
    """O modo limpo muda a apresentacao, nunca o quantitativo."""
    cfg_trechos.saida.modo_limpo = False
    _, completo = _abrir(cfg_trechos, prancha_trechos, tmp_path, "cheio.xlsx")
    cfg_trechos.saida.modo_limpo = True
    _, enxuto = _abrir(cfg_trechos, prancha_trechos, tmp_path, "limpo.xlsx")

    assert (sum(p.peso_com_perda_kg for p in enxuto.posicoes)
            == pytest.approx(sum(p.peso_com_perda_kg for p in completo.posicoes)))
    assert len(enxuto.posicoes) == len(completo.posicoes)


def test_dois_poligonos_na_mesma_layer_sao_UM_trecho(cfg_trechos, tmp_path):
    """Um trecho pode ser desenhado como duas regioes separadas (em L, ou
    partido por um vazio). As duas sao o MESMO trecho - renomear a segunda
    tiraria metade do trecho de qualquer conta feita por nome."""
    caminho = tmp_path / "trecho_partido.dxf"
    doc = ezdxf.new("R2010", setup=True)
    msp = doc.modelspace()
    for nome in ("TRECHO FT", "ARR_TXT", "ARR_BARRA"):
        doc.layers.add(nome)
    # duas regioes separadas, ambas na layer TRECHO FT
    msp.add_lwpolyline([(0, 0), (100, 0), (100, 100), (0, 100)], close=True,
                       dxfattribs={"layer": "TRECHO FT"})
    msp.add_lwpolyline([(300, 0), (400, 0), (400, 100), (300, 100)], close=True,
                       dxfattribs={"layer": "TRECHO FT"})
    for x, texto in ((50, "10N1-%%c10 C=400"), (350, "5N2-%%c8 C=200")):
        msp.add_line((x - 20, 50), (x + 20, 50), dxfattribs={"layer": "ARR_BARRA"})
        t = msp.add_text(texto, dxfattribs={"layer": "ARR_TXT", "height": 2.0})
        t.dxf.insert = (x, 50)
    doc.saveas(caminho)

    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(caminho, cfg_trechos)

    nomes = {a.nome for a in res.areas}
    assert nomes == {"TRECHO FT"}, f"trecho foi partido em {nomes}"
    assert len(res.areas) == 2, "os dois poligonos precisam existir"
    # as duas armaduras caem no mesmo trecho
    assert {p.area for p in res.posicoes} == {"TRECHO FT"}


# =============================================================================
# Contorno que nao fecha por ruido de coordenada
# =============================================================================

def test_contorno_com_erro_microscopico_e_costurado(cfg_trechos, tmp_path):
    """Um retangulo desenhado no CAD pode ter o canto com dois nos
    separados por 1e-13. Invisivel na tela, mas o polygonize exige
    coincidencia exata - sem costura o trecho inteiro desaparecia."""
    caminho = tmp_path / "canto_solto.dxf"
    doc = ezdxf.new("R2010", setup=True)
    msp = doc.modelspace()
    for nome in ("TRECHO A", "ARR_TXT", "ARR_BARRA"):
        doc.layers.add(nome)

    # Canto superior esquerdo com DOIS nos, separados por 1e-13:
    a1 = (192.4251402225746, 91.2372228739143)     # inicio da linha 1
    a2 = (192.4251402225745, 91.23722287391432)    # fim da linha 4
    b = (192.4251402225745, 78.19993486103942)
    c = (239.3379607080258, 78.19993486103942)
    d = (239.3379607080258, 91.23722287391433)
    for p, q in ((a1, b), (b, c), (c, d), (d, a2)):
        msp.add_line(p, q, dxfattribs={"layer": "TRECHO A"})

    msp.add_line((200, 82), (230, 82), dxfattribs={"layer": "ARR_BARRA"})
    t = msp.add_text("10N1-%%c10 C=400",
                     dxfattribs={"layer": "ARR_TXT", "height": 0.2})
    t.dxf.insert = (215, 82)
    doc.saveas(caminho)

    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(caminho, cfg_trechos)

    assert [a.nome for a in res.areas] == ["TRECHO A"], \
        "o contorno nao foi costurado e o trecho sumiu"
    assert {p.area for p in res.posicoes} == {"TRECHO A"}
    # e o programa avisa que precisou costurar
    assert "AREA_COSTURADA" in [i.tipo.name for i in res.inconsistencias]


def test_contorno_realmente_aberto_continua_sendo_erro(cfg_trechos, tmp_path):
    """A costura nao pode mascarar contorno aberto de verdade: um vao de
    centimetros e falha de desenho, nao ruido numerico."""
    caminho = tmp_path / "aberto_de_verdade.dxf"
    doc = ezdxf.new("R2010", setup=True)
    msp = doc.modelspace()
    for nome in ("TRECHO A", "ARR_TXT", "ARR_BARRA"):
        doc.layers.add(nome)
    # falta um lado inteiro
    for p, q in (((0, 0), (100, 0)), ((100, 0), (100, 100)),
                 ((100, 100), (0, 100))):
        msp.add_line(p, q, dxfattribs={"layer": "TRECHO A"})
    msp.add_line((20, 50), (80, 50), dxfattribs={"layer": "ARR_BARRA"})
    t = msp.add_text("10N1-%%c10 C=400",
                     dxfattribs={"layer": "ARR_TXT", "height": 2.0})
    t.dxf.insert = (50, 50)
    doc.saveas(caminho)

    configurar_log(cfg_trechos.log, tmp_path)
    res = processar(caminho, cfg_trechos)
    assert res.areas == []
    assert "AREA_NAO_FECHADA" in [i.tipo.name for i in res.inconsistencias]


# =============================================================================
# Recorte do desenho x prancha inteira
# =============================================================================

def test_recorte_nao_usa_a_tabela_do_projeto(cfg_trechos, tmp_path):
    """A tabela descreve o pavimento inteiro. Num recorte com poucas
    barras, ratear a tabela jogaria o pavimento todo em cima delas -
    uma barra desenhada virando centenas."""
    from extrator.tabela_mestre import LinhaMestre
    from extrator.pipeline import _usar_gabarito
    from extrator.modelos import ResultadoProcessamento, PosicaoArmadura

    gabarito = [LinhaMestre(posicao=str(i), bitola_mm=8.0, quantidade=10,
                            comprimento_total_cm=1000, peso_kg=39.5)
                for i in range(1, 101)]           # 100 posicoes na tabela
    desenhadas = [PosicaoArmadura(
        id="x", prancha="P", posicao=str(i), quantidade=1, bitola_mm=8.0,
        comprimento_unit_cm=100, espacamento_cm=None, texto_origem="")
        for i in range(1, 5)]                     # so 4 no desenho

    res = ResultadoProcessamento()
    cfg_trechos.calculo.fonte_quantidade = "auto"
    assert _usar_gabarito(cfg_trechos, gabarito, desenhadas, res, "P") is False
    assert "RECORTE_DETECTADO" in [i.tipo.name for i in res.inconsistencias]


def test_prancha_inteira_usa_a_tabela_do_projeto(cfg_trechos, tmp_path):
    from extrator.tabela_mestre import LinhaMestre
    from extrator.pipeline import _usar_gabarito
    from extrator.modelos import ResultadoProcessamento, PosicaoArmadura

    gabarito = [LinhaMestre(posicao=str(i), bitola_mm=8.0, quantidade=10,
                            comprimento_total_cm=1000, peso_kg=39.5)
                for i in range(1, 101)]
    desenhadas = [PosicaoArmadura(
        id="x", prancha="P", posicao=str(i), quantidade=1, bitola_mm=8.0,
        comprimento_unit_cm=100, espacamento_cm=None, texto_origem="")
        for i in range(1, 100)]                   # 99 das 100

    res = ResultadoProcessamento()
    cfg_trechos.calculo.fonte_quantidade = "auto"
    assert _usar_gabarito(cfg_trechos, gabarito, desenhadas, res, "P") is True


def test_o_rateio_nao_inventa_posicao_que_o_desenho_nao_tem(cfg_trechos):
    """Posicao da tabela sem barra desenhada nao entra no quantitativo -
    vira aviso. Era o que fazia um recorte reportar o pavimento inteiro."""
    from extrator.rateio import ratear_por_gabarito
    from extrator.tabela_mestre import LinhaMestre
    from extrator.modelos import ResultadoProcessamento, PosicaoArmadura

    gabarito = [
        LinhaMestre(posicao="1", bitola_mm=8.0, quantidade=10,
                    comprimento_total_cm=1000, peso_kg=39.5),
        LinhaMestre(posicao="99", bitola_mm=8.0, quantidade=500,
                    comprimento_total_cm=50000, peso_kg=1975.0),
    ]
    desenhadas = [PosicaoArmadura(
        id="x", prancha="P", posicao="1", quantidade=1, bitola_mm=8.0,
        comprimento_unit_cm=100, espacamento_cm=None, texto_origem="")]

    res = ResultadoProcessamento()
    novas = ratear_por_gabarito(desenhadas, gabarito, cfg_trechos, "P", res)
    assert [p.posicao for p in novas] == ["1"], "inventou a N99"
    assert novas[0].quantidade == 10          # a N1 recebeu a qtd da tabela
    assert "POSICAO_SEM_LOCALIZACAO" in [i.tipo.name for i in res.inconsistencias]


# =============================================================================
# MODO ENXUTO
# So a lista de corte e dobra de cada trecho e uma aba de conferencia.
# =============================================================================

def test_modo_enxuto_gera_so_as_abas_de_trecho_e_a_conferencia(
        cfg_trechos, prancha_trechos, tmp_path):
    cfg_trechos.saida.modo_enxuto = True
    wb, _ = _abrir(cfg_trechos, prancha_trechos, tmp_path)
    abas = wb.sheetnames

    assert "Obra - TRECHO 01" in abas and "Obra - TRECHO 02" in abas
    assert [a for a in abas if "CONFER" in a.upper()]
    # Nenhuma das abas de leitura pesada sobrevive.
    for indesejada in ("RESUMO GERAL", "COMPARACAO", "VERIFICACAO",
                       "INCONSIST"):
        assert not [a for a in abas if indesejada in a.upper()], indesejada
    # SEM_AREA tambem ganha aba: aco que nao caiu em trecho nenhum nao
    # pode sumir da planilha so porque o modo e enxuto.
    assert "Obra - SEM_AREA" in abas
    # Uma aba por trecho + a conferencia, e nada alem disso.
    assert len(abas) == 4


def test_modo_enxuto_MANTEM_a_lista_de_corte_e_dobra(cfg_trechos,
                                                     prancha_trechos, tmp_path):
    """Tirar aba nao pode tirar a medida da barra: sem posicao e sem
    comprimento unitario a planilha deixa de servir para cortar aco."""
    cfg_trechos.saida.modo_enxuto = True
    wb, _ = _abrir(cfg_trechos, prancha_trechos, tmp_path)
    texto = _texto_da_aba(wb["Obra - TRECHO 01"])
    assert "Posicao" in texto
    assert "Comprimento unitario (cm)" in texto


def test_modo_enxuto_tira_o_detalhamento_linha_a_linha(cfg_trechos,
                                                       prancha_trechos, tmp_path):
    cfg_trechos.saida.modo_enxuto = True
    cfg_trechos.saida.incluir_diagnostico = True
    wb, _ = _abrir(cfg_trechos, prancha_trechos, tmp_path)
    assert "DETALHAMENTO" not in _texto_da_aba(wb["Obra - TRECHO 01"]).upper()


def test_modo_enxuto_nao_altera_nenhum_numero(cfg_trechos, prancha_trechos,
                                              tmp_path):
    """A planilha muda de forma, nunca de conteudo."""
    cfg_trechos.saida.modo_enxuto = False
    _, completo = _abrir(cfg_trechos, prancha_trechos, tmp_path, "a.xlsx")
    cfg_trechos.saida.modo_enxuto = True
    _, enxuto = _abrir(cfg_trechos, prancha_trechos, tmp_path, "b.xlsx")

    peso = lambda r: round(sum(p.peso_com_perda_kg for p in r.posicoes), 6)
    assert peso(completo) == peso(enxuto)
    assert len(completo.posicoes) == len(enxuto.posicoes)


def test_modo_enxuto_sem_tabela_mestre_diz_que_nao_ha_o_que_conferir(
        cfg_trechos, prancha_trechos, tmp_path):
    """O DXF sintetico nao tem tabela de aco desenhada. A aba nao pode
    fingir que conferiu: tem de dizer que faltou o gabarito."""
    cfg_trechos.saida.modo_enxuto = True
    wb, res = _abrir(cfg_trechos, prancha_trechos, tmp_path)
    assert not res.tabela_mestre
    aba = [a for a in wb.sheetnames if "CONFER" in a.upper()][0]
    texto = _texto_da_aba(wb[aba]).lower()
    assert "nenhuma tabela" in texto


def test_modo_enxuto_vale_mais_que_o_modo_limpo(cfg_trechos, prancha_trechos,
                                                tmp_path):
    """Os dois ligados nao podem gerar uma planilha meio-termo."""
    cfg_trechos.saida.modo_limpo = True
    cfg_trechos.saida.modo_enxuto = True
    wb, _ = _abrir(cfg_trechos, prancha_trechos, tmp_path)
    assert not [a for a in wb.sheetnames if "RESUMO GERAL" in a.upper()]
    assert [a for a in wb.sheetnames if "CONFER" in a.upper()]
